# -*- coding: utf-8 -*-
import base64
import io
from odoo import _, api, fields, models
from odoo.exceptions import UserError


class StockExpiryExport(models.TransientModel):
    _name = 'stock.expiry.export'
    _description = 'Export Expiry Report Wizard'

    export_format = fields.Selection([
        ('pdf', 'PDF Report'),
        ('xlsx', 'Excel (XLSX)'),
    ], string='Export Format', default='pdf', required=True)
    status_filter = fields.Selection([
        ('all', 'All Lots with Expiry Date'),
        ('red', 'RED (Critical) only'),
        ('yellow_red', 'YELLOW + RED'),
        ('expired', 'Expired only'),
    ], string='Filter by Status', default='all', required=True)
    date_from = fields.Date(string='Expiry Date From')
    date_to = fields.Date(string='Expiry Date To')

    export_file = fields.Binary(string='Download', readonly=True)
    export_filename = fields.Char(string='Filename', readonly=True)

    def _get_lots(self):
        domain = [('expiration_date', '!=', False)]
        if self.status_filter == 'red':
            domain += [('expiry_status', '=', 'red')]
        elif self.status_filter == 'yellow_red':
            domain += [('expiry_status', 'in', ['yellow', 'red'])]
        elif self.status_filter == 'expired':
            domain += [('expiry_status', '=', 'expired')]
        if self.date_from:
            domain += [('expiration_date', '>=', fields.Datetime.to_datetime(self.date_from))]
        if self.date_to:
            domain += [('expiration_date', '<=', fields.Datetime.to_datetime(self.date_to))]
        return self.env['stock.lot'].search(domain, order='expiration_date asc')

    def action_export(self):
        self.ensure_one()
        lots = self._get_lots()
        if not lots:
            raise UserError(_('No lots found matching the selected criteria.'))

        if self.export_format == 'pdf':
            return self._export_pdf(lots)
        return self._export_xlsx(lots)

    def _export_pdf(self, lots):
        report = self.env.ref('stock_expiry_alert_manager.action_report_expiry_lots')
        pdf_content, _ = self.env['ir.actions.report']._render_qweb_pdf(
            report, lots.ids)
        self.export_file = base64.b64encode(pdf_content)
        self.export_filename = 'expiry_report.pdf'
        return {
            'type': 'ir.actions.act_window',
            'res_model': 'stock.expiry.export',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }

    def _export_xlsx(self, lots):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            raise UserError(_('openpyxl is required for XLSX export. Install it with: pip install openpyxl'))

        STATUS_COLORS = {
            'green': 'C6EFCE',
            'yellow': 'FFEB9C',
            'red': 'FFC7CE',
            'expired': 'D9D9D9',
        }

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Expiry Report'

        headers = ['Lot/Serial', 'Product', 'Status', 'Days to Expiry',
                   'Expiry Date', 'Qty on Hand', 'Estimated Cost', 'Location']
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(fill_type='solid', fgColor='2E4057')

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        for row, lot in enumerate(lots, 2):
            status = lot.expiry_status or 'green'
            fill = PatternFill(fill_type='solid', fgColor=STATUS_COLORS.get(status, 'FFFFFF'))
            quants = self.env['stock.quant'].search([
                ('lot_id', '=', lot.id),
                ('quantity', '>', 0),
                ('location_id.usage', '=', 'internal'),
            ])
            location_names = ', '.join(quants.mapped('location_id.complete_name'))

            values = [
                lot.name,
                lot.product_id.display_name,
                dict(lot._fields['expiry_status'].selection).get(status, status).upper(),
                lot.days_to_expiry,
                lot.expiration_date.strftime('%Y-%m-%d') if lot.expiration_date else '',
                lot.product_qty,
                lot.estimated_cost,
                location_names,
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.fill = fill

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        self.export_file = base64.b64encode(output.read())
        self.export_filename = 'expiry_report.xlsx'

        return {
            'type': 'ir.actions.act_window',
            'res_model': 'stock.expiry.export',
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
        }
